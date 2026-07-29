"""
변우석 아카이브 - 폴더 기반 메타데이터 엑셀 자동 생성 프로그램
================================================================

폴더 구조 규칙 (5단계, 루트 폴더 기준):
    루트/연도/대분류/중분류/출처/비고/파일

    예) 루트/2024/앰버서더/프라다/매거진/보그_화보/photo001.jpg

- 중간 단계 폴더가 없으면 자동으로 "N/A"로 채워집니다.
  (예: 루트/2024/앰버서더/프라다/photo001.jpg 처럼 출처, 비고 폴더가
   없으면 출처=N/A, 비고=N/A 로 자동 처리)
- 사진(jpg/jpeg/png/gif/webp/bmp/heic)과 영상(mp4/mov/avi/mkv/wmv) 둘 다 처리합니다.
- 영상은 1초 지점 프레임을 추출해 썸네일로 사용합니다. 추출이 안 되면 N/A로 표시합니다.

사용법:
    python archive_excel_generator.py --root "정리할 폴더 경로" --output "결과.xlsx"

    예) python archive_excel_generator.py --root "/Users/me/우석사진" --output "우석_아카이브_대장.xlsx"
"""

import argparse
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

try:
    import cv2
    HAS_CV2 = True
except ImportError:
    HAS_CV2 = False

PHOTO_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".heic", ".tiff"}
VIDEO_EXTS = {".mp4", ".mov", ".avi", ".mkv", ".wmv", ".m4v"}

HEADERS = ["파일명", "연도", "대분류", "중분류", "출처", "비고", "유형", "썸네일"]
FIELD_ORDER = ["연도", "대분류", "중분류", "출처", "비고"]  # 폴더 depth 순서

THUMB_SIZE = 100  # 썸네일 한 변 픽셀 크기 (셀에 삽입되는 이미지)
ROW_HEIGHT = 80    # 썸네일이 들어가는 행 높이 (포인트)


def classify_file(path: Path) -> str:
    """확장자를 보고 사진/영상/기타를 구분."""
    ext = path.suffix.lower()
    if ext in PHOTO_EXTS:
        return "사진"
    if ext in VIDEO_EXTS:
        return "영상"
    return "기타"


def parse_folder_fields(rel_parts):
    """
    루트 기준 상대 경로의 폴더 부분(rel_parts, 파일명 제외)을 받아서
    연도/대분류/중분류/출처/비고 5개 필드로 매핑.

    - 폴더 depth가 5보다 적으면: 앞에서부터 채우고 나머지는 N/A
    - 폴더 depth가 5보다 많으면: 5번째 이후 폴더명은 전부 '비고'에 이어붙임
    """
    fields = {}
    n = len(rel_parts)
    if n >= 5:
        fields["연도"] = rel_parts[0]
        fields["대분류"] = rel_parts[1]
        fields["중분류"] = rel_parts[2]
        fields["출처"] = rel_parts[3]
        # 5번째 이후 폴더가 더 있으면 비고에 이어붙임
        extra = rel_parts[4:]
        fields["비고"] = "_".join(extra) if extra else "N/A"
    else:
        for i, key in enumerate(FIELD_ORDER):
            fields[key] = rel_parts[i] if i < n else "N/A"

    for key in FIELD_ORDER:
        if not fields.get(key):
            fields[key] = "N/A"
    return fields


def make_photo_thumbnail(src_path: Path, tmp_dir: Path):
    """사진 파일을 축소해서 임시 썸네일 파일로 저장. 실패하면 None."""
    try:
        with PILImage.open(src_path) as img:
            img = img.convert("RGB")
            img.thumbnail((THUMB_SIZE, THUMB_SIZE))
            out_path = tmp_dir / f"thumb_{abs(hash(str(src_path)))}.jpg"
            img.save(out_path, "JPEG", quality=85)
            return out_path
    except Exception:
        return None


def make_video_thumbnail(src_path: Path, tmp_dir: Path):
    """영상 파일의 1초 지점 프레임을 추출해서 임시 썸네일로 저장. 실패하면 None."""
    if not HAS_CV2:
        return None
    try:
        cap = cv2.VideoCapture(str(src_path))
        fps = cap.get(cv2.CAP_PROP_FPS) or 0
        frame_no = int(fps * 1) if fps and fps > 0 else 0
        cap.set(cv2.CAP_PROP_POS_FRAMES, frame_no)
        ok, frame = cap.read()
        if not ok:
            cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
            ok, frame = cap.read()
        cap.release()
        if not ok:
            return None
        frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        img = PILImage.fromarray(frame)
        img.thumbnail((THUMB_SIZE, THUMB_SIZE))
        out_path = tmp_dir / f"thumb_{abs(hash(str(src_path)))}.jpg"
        img.save(out_path, "JPEG", quality=85)
        return out_path
    except Exception:
        return None


def collect_files(root: Path):
    """루트 폴더를 재귀적으로 훑어서 (파일경로, 상대 폴더 파츠) 리스트 반환."""
    results = []
    for dirpath, _dirnames, filenames in os.walk(root):
        for fname in filenames:
            if fname.startswith("."):
                continue
            fpath = Path(dirpath) / fname
            ext = fpath.suffix.lower()
            if ext not in PHOTO_EXTS and ext not in VIDEO_EXTS:
                continue
            rel = fpath.relative_to(root)
            rel_parts = list(rel.parts[:-1])  # 파일명 제외한 폴더 부분
            results.append((fpath, rel_parts))
    results.sort(key=lambda x: str(x[0]).lower())
    return results


def safe_sheet_name(name: str, used_names: set) -> str:
    """엑셀 시트명 규칙(31자 제한, 특수문자 금지, 중복 금지)에 맞게 정리."""
    invalid = set('\\/*?:[]')
    cleaned = "".join(c for c in str(name) if c not in invalid).strip() or "미분류"
    cleaned = cleaned[:31]
    base = cleaned
    n = 1
    while cleaned in used_names:
        suffix = f"_{n}"
        cleaned = base[: 31 - len(suffix)] + suffix
        n += 1
    used_names.add(cleaned)
    return cleaned


def write_sheet(ws, rows, tmp_dir):
    """한 시트(탭)에 헤더 + 데이터 행 + 썸네일을 채워 넣는다. 반환값: 썸네일 실패 개수."""
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    body_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    col_widths = [30, 8, 14, 20, 12, 20, 8, 16]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row_idx = 2
    no_thumb_count = 0

    for fpath, rel_parts in rows:
        fields = parse_folder_fields(rel_parts)
        ftype = classify_file(fpath)

        ws.cell(row=row_idx, column=1, value=fpath.name)
        ws.cell(row=row_idx, column=2, value=fields["연도"])
        ws.cell(row=row_idx, column=3, value=fields["대분류"])
        ws.cell(row=row_idx, column=4, value=fields["중분류"])
        ws.cell(row=row_idx, column=5, value=fields["출처"])
        ws.cell(row=row_idx, column=6, value=fields["비고"])
        ws.cell(row=row_idx, column=7, value=ftype)

        for col in range(1, 8):
            c = ws.cell(row=row_idx, column=col)
            c.font = body_font
            c.border = border
            if col != 1:
                c.alignment = center

        thumb_path = None
        if ftype == "사진":
            thumb_path = make_photo_thumbnail(fpath, tmp_dir)
        elif ftype == "영상":
            thumb_path = make_video_thumbnail(fpath, tmp_dir)

        if thumb_path and thumb_path.exists():
            xl_img = XLImage(str(thumb_path))
            xl_img.width = THUMB_SIZE
            xl_img.height = THUMB_SIZE
            anchor_cell = f"{get_column_letter(8)}{row_idx}"
            ws.add_image(xl_img, anchor_cell)
        else:
            ws.cell(row=row_idx, column=8, value="N/A")
            ws.cell(row=row_idx, column=8).font = body_font
            ws.cell(row=row_idx, column=8).alignment = center
            ws.cell(row=row_idx, column=8).border = border
            no_thumb_count += 1

        ws.row_dimensions[row_idx].height = ROW_HEIGHT
        row_idx += 1

    ws.freeze_panes = "A2"
    return no_thumb_count


def build_excel(root: Path, output_path: Path):
    tmp_dir = root.parent / "._thumb_tmp"
    tmp_dir.mkdir(exist_ok=True)

    files = collect_files(root)
    if not files:
        print(f"경고: '{root}' 아래에서 사진/영상 파일을 찾지 못했습니다.")

    # 연도별로 그룹핑 (연도 폴더가 없으면 N/A 그룹으로)
    groups = {}
    for fpath, rel_parts in files:
        year = rel_parts[0] if rel_parts else "N/A"
        groups.setdefault(year, []).append((fpath, rel_parts))

    # 연도 내림차순 정렬 (최신 연도가 앞 탭), N/A는 맨 뒤로
    sorted_years = sorted([y for y in groups if y != "N/A"], reverse=True) + (
        ["N/A"] if "N/A" in groups else []
    )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 기본 빈 시트 제거

    used_names = set()
    total_no_thumb = 0
    for year in sorted_years:
        sheet_name = safe_sheet_name(year, used_names)
        ws = wb.create_sheet(title=sheet_name)
        no_thumb = write_sheet(ws, groups[year], tmp_dir)
        total_no_thumb += no_thumb

    wb.save(output_path)

    print(f"완료: {len(files)}개 파일 처리, {len(sorted_years)}개 연도 탭 생성 -> {output_path}")
    for year in sorted_years:
        print(f"  - {year}: {len(groups[year])}개")
    if total_no_thumb:
        print(f"  (썸네일 생성 실패/영상 미지원으로 N/A 처리된 항목: {total_no_thumb}개)")


def main():
    parser = argparse.ArgumentParser(description="폴더 구조 기반 아카이브 메타데이터 엑셀 생성기")
    parser.add_argument("--root", required=True, help="정리할 최상위 루트 폴더 경로")
    parser.add_argument("--output", default="아카이브_대장.xlsx", help="결과 엑셀 파일 경로")
    args = parser.parse_args()

    root = Path(args.root).expanduser().resolve()
    if not root.exists() or not root.is_dir():
        print(f"오류: '{root}' 폴더를 찾을 수 없습니다.")
        sys.exit(1)

    output_path = Path(args.output).expanduser().resolve()
    build_excel(root, output_path)


if __name__ == "__main__":
    main()
